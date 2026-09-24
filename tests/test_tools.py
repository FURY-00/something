import threading

import pytest

from jarvis.config import load_config
from jarvis.context import ToolContext
from jarvis.llm import ChatResult
from jarvis.memory import Memory
from jarvis.tools import ToolRegistry, _REGISTRY, tool
from jarvis.tools import system, web, writing
from jarvis.tools.control import parse_keys
from jarvis.wakeword import strip_wake_word
from jarvis.assistant import is_yes


@pytest.fixture
def ctx(tmp_path):
    cfg = load_config()
    cfg.paths.documents = tmp_path / "docs"
    return ToolContext(cfg, memory=Memory(tmp_path / "mem.json"))


def test_all_builtin_tools_have_valid_schemas(ctx):
    registry = ToolRegistry(ctx)
    for schema in registry.schemas():
        fn = schema["function"]
        assert fn["description"]
        props = fn["parameters"]["properties"]
        assert set(fn["parameters"]["required"]) <= set(props)
        for spec in props.values():
            assert spec["type"] in {"string", "integer", "number", "boolean"}
            assert spec.get("description")
    # Optional tools stay hidden until enabled in config.
    assert "run_shell_command" not in registry.names()
    assert "lookup_encyclopedia" not in registry.names()
    assert "look_at_screen" not in registry.names()


def test_optional_tools_appear_when_enabled(ctx):
    ctx.config.safety.allow_shell = True
    ctx.config.knowledge.kiwix_url = "http://localhost:8080"
    ctx.config.llm.vision_model = "llava:7b"
    names = ToolRegistry(ctx).names()
    assert {"run_shell_command", "lookup_encyclopedia", "look_at_screen"} <= set(names)


def test_registry_coerces_and_validates_arguments(ctx):
    seen = {}

    @tool("t", {
        "n": {"type": "integer", "description": "n"},
        "flag": {"type": "boolean", "description": "f"},
        "mode": {"type": "string", "enum": ["a", "b"], "description": "m"},
    }, name="coerce_test")
    def coerce_test(ctx, n, flag=False, mode="a"):
        seen.update(n=n, flag=flag, mode=mode)
        return None

    registry = ToolRegistry(ctx, {"coerce_test": _REGISTRY.pop("coerce_test")})
    assert registry.call("coerce_test", {"n": "3", "flag": "yes", "junk": 1}) == "Done."
    assert seen == {"n": 3, "flag": True, "mode": "a"}
    assert "missing required" in registry.call("coerce_test", {})
    assert "must be one of" in registry.call("coerce_test", {"n": 1, "mode": "z"})
    assert "no tool called" in registry.call("nope", {})


def test_tool_errors_are_reported_not_raised(ctx):
    @tool("boom", name="boom_test")
    def boom_test(ctx):
        raise RuntimeError("kaput")

    registry = ToolRegistry(ctx, {"boom_test": _REGISTRY.pop("boom_test")})
    assert registry.call("boom_test", {}) == "Error while running boom_test: kaput"


def test_undocumented_params_are_rejected():
    with pytest.raises(ValueError):
        @tool("bad")
        def bad(ctx, x):
            pass


def test_confirm_respects_config(ctx):
    asked = []
    ctx._ask_yes_no = lambda q: asked.append(q) or False
    assert ctx.confirm("Shut down?") is False
    ctx.config.safety.confirm_dangerous = False
    assert ctx.confirm("Shut down?") is True
    assert asked == ["Shut down?"]


def test_power_actions_ask_first(ctx, monkeypatch):
    launched = []
    monkeypatch.setattr(system.subprocess, "Popen", lambda cmd: launched.append(cmd))
    ctx._ask_yes_no = lambda q: False
    assert system.system_power(ctx, "shutdown") == "Cancelled by the user."
    assert launched == []
    assert system.system_power(ctx, "lock") == "Okay, lock."
    assert len(launched) == 1


def test_shell_always_asks_even_without_confirmations(ctx):
    ctx.config.safety.confirm_dangerous = False
    ctx._ask_yes_no = lambda q: False
    assert system.run_shell_command(ctx, "echo hi") == "Cancelled by the user."
    ctx._ask_yes_no = lambda q: True
    assert "hi" in system.run_shell_command(ctx, "echo hi")


def test_timer_announces(ctx):
    done = threading.Event()
    heard = []
    ctx._announce = lambda text: (heard.append(text), done.set())
    assert system.set_timer(ctx, 0.001, "stretch") == "Timer set for 1 second."
    assert done.wait(3)
    assert heard == ["Sir, this is your reminder: stretch."]


def test_human_duration():
    assert system._human_duration(90) == "1 minute and 30 seconds"
    assert system._human_duration(7200) == "2 hours"


def test_open_website_normalizes(ctx, monkeypatch):
    opened = []
    monkeypatch.setattr(web.webbrowser, "open", opened.append)
    web.open_website(ctx, "YouTube")
    web.open_website(ctx, "example.org")
    web.open_website(ctx, "stack overflow")
    web.search_web(ctx, "cute cats", "youtube")
    assert opened == [
        "https://www.youtube.com",
        "https://example.org",
        "https://stackoverflow.com",
        "https://www.youtube.com/results?search_query=cute+cats",
    ]


def test_write_document_saves_word_file(ctx, monkeypatch):
    class LLM:
        def chat(self, messages, **kw):
            assert "abstract" in messages[-1]["content"]
            return ChatResult(content="<think>plan</think># Solar Power: An Abstract\n\n"
                                      "Solar power is **clean**.\n\n- cheap\n- *abundant*")

    ctx.llm = LLM()
    opened, said = [], []
    ctx._say = said.append
    monkeypatch.setattr(writing, "open_with_default_app", opened.append)
    result = writing.write_document(ctx, "a short abstract on solar power")
    assert "Solar Power: An Abstract" in result
    path = ctx.config.paths.documents / "Solar Power An Abstract.docx"
    assert opened == [str(path)] and path.exists()
    from docx import Document

    doc = Document(str(path))
    texts = [p.text for p in doc.paragraphs]
    assert texts == ["Solar Power: An Abstract", "Solar power is clean.", "cheap", "abundant"]
    assert doc.paragraphs[1].runs[1].bold
    # A second document with the same title doesn't overwrite the first.
    writing.write_document(ctx, "another abstract")
    assert (ctx.config.paths.documents / "Solar Power An Abstract (2).docx").exists()
    assert said and "On it" in said[0]


def test_markdown_to_plain():
    assert writing.markdown_to_plain("# Title\n\n* **a** b\n```\ncode\n```") == "Title\n\n- a b\ncode\n"


def test_memory_persists(tmp_path):
    mem = Memory(tmp_path / "m.json")
    mem.add("My birthday is 10 May.")
    mem.add("my birthday is 10 may.")
    again = Memory(tmp_path / "m.json")
    assert again.facts == ["My birthday is 10 May."]
    assert again.forget("birthday") == ["My birthday is 10 May."]
    assert Memory(tmp_path / "m.json").facts == []


def test_parse_keys():
    assert parse_keys("Control + C") == ["ctrl", "c"]
    assert parse_keys("alt+tab") == ["alt", "tab"]
    assert parse_keys("page down") == ["pagedown"]


def test_strip_wake_word():
    names = ["jarvis", "travis"]
    assert strip_wake_word("Hey Jarvis, open YouTube.", names) == "open YouTube"
    assert strip_wake_word("What time is it, Jarvis?", names) == "What time is it"
    assert strip_wake_word("Jarvis.", names) == ""
    assert strip_wake_word("Open YouTube", names) is None


def test_is_yes():
    assert is_yes("Yes please")
    assert is_yes("yeah, go ahead")
    assert not is_yes("No, don't")
    assert not is_yes("")
    assert not is_yes("wait, no")
