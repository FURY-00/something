"""Documents -> text/images, the website builder and the program writer."""

import io

import pytest

from jarvis.config import load_config
from jarvis.context import ToolContext
from jarvis.documents import extract, find_documents
from jarvis.llm import ChatResult
from jarvis.skills import web


def noisy_png(size=64) -> bytes:
    from PIL import Image

    img = Image.effect_noise((size, size), 80).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def simple_pdf(lines) -> bytes:
    """A tiny valid one-page PDF with text, written by hand."""
    stream = "BT /F1 18 Tf 72 720 Td " + " ".join(f"({t}) Tj 0 -24 Td" for t in lines) + " ET"
    objs = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


@pytest.fixture
def docs(tmp_path):
    from docx import Document
    from pptx import Presentation
    from pptx.util import Inches

    png = tmp_path / "chart.png"
    png.write_bytes(noisy_png())
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Carnot cycle"
    slide.placeholders[1].text = "Efficiency = 1 - Tc/Th"
    slide.shapes.add_picture(str(png), Inches(1), Inches(3))
    slide.notes_slide.notes_text_frame.text = "Mention reversibility"
    prs.save(tmp_path / "lecture.pptx")
    doc = Document()
    doc.add_heading("Heat engines", 1)
    doc.add_paragraph("A heat engine converts heat into work.")
    doc.add_picture(str(png))
    doc.save(tmp_path / "notes.docx")
    (tmp_path / "summary.pdf").write_bytes(simple_pdf(["Entropy never decreases", "in an isolated system"]))
    return tmp_path


def test_extracts_text_and_images(docs, tmp_path):
    assets = tmp_path / "assets"
    pptx = extract(docs / "lecture.pptx", assets)
    assert "## Slide 1: Carnot cycle" in pptx.text and "Efficiency = 1 - Tc/Th" in pptx.text
    assert "(Speaker notes: Mention reversibility)" in pptx.text
    assert len(pptx.images) == 1 and pptx.images[0].exists()
    word = extract(docs / "notes.docx", assets)
    assert "# Heat engines" in word.text and "converts heat into work" in word.text
    assert len(word.images) == 1
    pdf = extract(docs / "summary.pdf")
    assert "Entropy never decreases" in pdf.text and "## Page 1" in pdf.text
    found = find_documents([str(docs)])
    assert {p.name for p in found} >= {"lecture.pptx", "notes.docx", "summary.pdf"}


class CodeLLM:
    def __init__(self, *replies):
        self.replies = list(replies)
        self.prompts = []

    def chat(self, messages, **kw):
        self.prompts.append(messages[-1]["content"])
        return ChatResult(content=self.replies.pop(0))


@pytest.fixture
def ctx(tmp_path, monkeypatch):
    cfg = load_config()
    cfg.paths.projects = tmp_path / "projects"
    opened = []
    monkeypatch.setattr("jarvis.tools.system.open_with_default_app", opened.append)
    c = ToolContext(cfg, ask_yes_no=lambda q: True)
    c.opened = opened
    return c


def test_builds_a_website_from_documents(docs, ctx):
    page = ("### FILE: index.html\n```html\n<!doctype html><html><head><title>Thermo</title></head><body>"
            "<h1>Carnot cycle</h1><img src=\"assets/lecture-s1-1.png\" alt=\"chart\"><script>"
            "document.querySelector('h1').onclick=()=>alert('hi')</script></body></html>\n```")
    ctx.code_llm = CodeLLM(page)
    result = web.build_site(ctx, "An interactive study site with a quiz", f"{docs / 'lecture.pptx'}", "Thermo Notes")
    site = ctx.config.paths.projects / "websites" / "thermo-notes"
    assert "index.html" in result and (site / "index.html").exists()
    assert (site / "assets" / "lecture-s1-1.png").exists()
    assert "assets/lecture-s1-1.png" in ctx.code_llm.prompts[0]  # the model was told about the image
    assert "Efficiency = 1 - Tc/Th" in ctx.code_llm.prompts[0]   # and given the slide text
    assert ctx.opened == [str(site / "index.html")]


def test_broken_links_are_sent_back_to_be_fixed(ctx):
    bad = "### FILE: index.html\n```html\n<html><body><p>x</p><img src=\"missing.png\"></body></html>\n```"
    good = "### FILE: index.html\n```html\n<html><body><p>x</p><p>fixed</p></body></html>\n```"
    ctx.code_llm = CodeLLM(bad, good)
    result = web.build_site(ctx, "a page", "", "demo")
    assert "Remaining problems" not in result
    assert "missing.png, which doesn't exist" in ctx.code_llm.prompts[1]


def test_editing_keeps_a_backup(ctx):
    site = ctx.config.paths.projects / "websites" / "demo"
    site.mkdir(parents=True)
    (site / "index.html").write_text("<html><body><h1>Old</h1><p>a</p></body></html>")
    ctx.code_llm = CodeLLM("### FILE: index.html\n```html\n<html><body><h1>New</h1><p>a</p></body></html>\n```")
    web.edit_project(ctx, web.find_project(ctx, "", "websites"), "rename the heading", site=True)
    assert "New" in (site / "index.html").read_text()
    backups = list((site / ".history").rglob("index.html"))
    assert backups and "Old" in backups[0].read_text()
    assert "<h1>Old</h1>" in ctx.code_llm.prompts[0]  # the model saw the current file


def test_python_program_is_run_and_fixed(ctx):
    broken = "### FILE: main.py\n```python\nprint(1/0)\n```"
    fixed = "### FILE: main.py\n```python\nprint('area =', 3.14159 * 2 ** 2)\n```"
    ctx.code_llm = CodeLLM(broken, fixed)
    result = web.write_program(ctx, "print the area of a circle of radius 2", "python", "circle", run=True)
    assert "area = 12.56636" in result
    assert "ZeroDivisionError" in ctx.code_llm.prompts[1]


def test_generated_paths_cannot_escape_the_project():
    files = web.parse_files("### FILE: ../../evil.py\n```python\nx\n```\n### FILE: ok.py\n```python\ny\n```", "main.py")
    assert list(files) == ["ok.py"]


def test_content_is_condensed_for_small_models(ctx):
    ctx.code_llm = CodeLLM(*["notes"] * 20)
    said = []
    out = web.condense("word " * 20000, 30000, ctx.code_llm, said.append)
    assert len(out) <= 30000 and said
    assert web.condense("short", 1000, ctx.code_llm, said.append) == "short"


def test_unknown_project_name_is_not_confused_with_another(ctx):
    base = ctx.config.paths.projects / "websites"
    (base / "thermo-notes").mkdir(parents=True)
    assert web.find_project(ctx, "", "websites").name == "thermo-notes"      # most recent
    assert web.find_project(ctx, "thermo", "websites").name == "thermo-notes"  # partial name
    missing = web.find_project(ctx, "fluid mechanics", "websites")
    assert missing.name == "fluid-mechanics" and not missing.exists()
    assert web.edit_project(ctx, missing, "x", site=True).startswith("Error")
