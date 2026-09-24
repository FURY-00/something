"""Read text (and pictures) out of PDF, PowerPoint, Word, text and HTML files."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED = {".pdf", ".pptx", ".docx", ".txt", ".md", ".html", ".htm", ".csv", ".json", ".py",
             ".js", ".css"}
MIN_IMAGE_BYTES = 6000  # skip tiny logos, bullets and spacer images


@dataclass
class Extracted:
    path: Path
    kind: str
    text: str
    images: list[Path] = field(default_factory=list)

    @property
    def words(self) -> int:
        return len(self.text.split())


def find_documents(paths: list[str]) -> list[Path]:
    """Expand files and folders into the supported documents inside them."""
    out: list[Path] = []
    for raw in paths:
        p = Path(raw.strip().strip('"')).expanduser()
        if p.is_dir():
            out += sorted(f for f in p.rglob("*") if f.suffix.lower() in SUPPORTED and f.is_file())
        elif p.exists():
            out.append(p)
        else:
            raise FileNotFoundError(f"{raw} doesn't exist")
    return out


def extract(path: str | Path, assets_dir: Path | None = None, max_images: int = 40) -> Extracted:
    path = Path(path).expanduser()
    suffix = path.suffix.lower()
    saver = _ImageSaver(assets_dir, path.stem, max_images)
    if suffix == ".pdf":
        text = _pdf(path, saver)
    elif suffix == ".pptx":
        text = _pptx(path, saver)
    elif suffix == ".docx":
        text = _docx(path, saver)
    elif suffix in (".html", ".htm"):
        text = _html(path.read_text(encoding="utf-8", errors="replace"))
    else:
        text = path.read_text(encoding="utf-8", errors="replace")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return Extracted(path, suffix.lstrip("."), text, saver.saved)


class _ImageSaver:
    def __init__(self, folder: Path | None, stem: str, limit: int) -> None:
        self.folder, self.stem, self.limit = folder, re.sub(r"[^\w-]+", "-", stem).strip("-"), limit
        self.saved: list[Path] = []

    def save(self, data: bytes, ext: str, label: str) -> str | None:
        if self.folder is None or len(data) < MIN_IMAGE_BYTES or len(self.saved) >= self.limit:
            return None
        ext = ext.lower().lstrip(".") or "png"
        if ext not in {"png", "jpg", "jpeg", "gif", "webp", "svg", "bmp"}:
            return None
        self.folder.mkdir(parents=True, exist_ok=True)
        target = self.folder / f"{self.stem}-{label}-{len(self.saved) + 1}.{ext}"
        target.write_bytes(data)
        self.saved.append(target)
        return target.name


def _pdf(path: Path, saver: _ImageSaver) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        parts.append(f"## Page {i}\n{page.extract_text() or ''}")
        try:
            for img in page.images:
                name = saver.save(img.data, Path(img.name).suffix, f"p{i}")
                if name:
                    parts.append(f"[image: {name}]")
        except Exception:  # noqa: BLE001 - some PDFs have unreadable image streams
            pass
    return "\n".join(parts)


def _pptx(path: Path, saver: _ImageSaver) -> str:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    parts = []

    def walk(shapes, slide_no, out):
        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, slide_no, out)
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text.strip())
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text.strip() for c in row.cells))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                name = saver.save(shape.image.blob, shape.image.ext, f"s{slide_no}")
                if name:
                    out.append(f"[image: {name}]")

    for i, slide in enumerate(prs.slides, 1):
        title = slide.shapes.title.text.strip() if slide.shapes.title is not None else ""
        body: list[str] = []
        walk(slide.shapes, i, body)
        if title and body and body[0] == title:
            body = body[1:]
        parts.append(f"## Slide {i}: {title}".rstrip(": ") + "\n" + "\n".join(body))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"(Speaker notes: {notes})")
    return "\n\n".join(parts)


def _docx(path: Path, saver: _ImageSaver) -> str:
    from docx import Document

    doc = Document(str(path))
    parts = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        level = re.search(r"heading (\d)", style)
        parts.append(("#" * int(level.group(1)) + " " + text) if level else
                     ("# " + text if style == "title" else text))
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    for rel in doc.part.rels.values():
        if "image" in rel.reltype:
            part = rel.target_part
            name = saver.save(part.blob, Path(str(part.partname)).suffix, "img")
            if name:
                parts.append(f"[image: {name}]")
    return "\n".join(parts)


class _Text(HTMLParser):
    SKIP = {"script", "style", "noscript"}

    def __init__(self) -> None:
        super().__init__()
        self.out: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in self.SKIP:
            self._skip += 1
        elif tag in ("p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"):
            self.out.append("\n")

    def handle_endtag(self, tag):
        if tag in self.SKIP and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.out.append(data)


def _html(html: str) -> str:
    parser = _Text()
    parser.feed(html)
    return re.sub(r"[ \t]+", " ", "".join(parser.out))
