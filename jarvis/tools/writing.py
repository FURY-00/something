"""Writing documents: abstracts, essays, emails, reports, letters, code..."""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from ..llm import strip_think
from . import tool
from .control import _paste
from .system import open_with_default_app

WRITER_PROMPT = (
    "You are an expert writer. Write exactly what the user asks for, at the "
    "requested length and style, with correct facts. Output only the finished "
    "document in Markdown: a '# Title' line first, then the content. No preamble, "
    "no notes, no closing remarks."
)


@tool(
    "Write a longer piece of text (an abstract, essay, email, letter, report, "
    "summary, story, poem, notes, code) and save it as a document, or type it "
    "into the window that is currently active. Use this whenever the user wants "
    "something written, instead of speaking the whole text.",
    {
        "request": {
            "type": "string",
            "description": "Everything about what to write: topic, type, length, tone, "
            "audience and any details the user mentioned",
        },
        "destination": {
            "type": "string",
            "enum": ["word", "text", "type"],
            "description": "word = save a Word .docx and open it (default); text = save "
            "a plain text/markdown file and open it; type = type it at the cursor in "
            "the active window (e.g. an open Google Doc or email)",
        },
        "filename": {"type": "string", "description": "Optional file name, without extension"},
    },
)
def write_document(ctx, request: str, destination: str = "word", filename: str = "") -> str:
    ctx.say("On it. Writing takes a moment.")
    result = ctx.llm.chat(
        [{"role": "system", "content": WRITER_PROMPT}, {"role": "user", "content": request}],
        temperature=0.7,
    )
    markdown = strip_think(result.content)
    if not markdown:
        return "Error: the language model returned nothing."
    title = _title(markdown) or filename or request[:50]
    words = len(markdown.split())

    if destination == "type":
        _paste(markdown_to_plain(markdown))
        return f"Typed the {words}-word text '{title}' into the active window."

    folder: Path = ctx.config.paths.documents
    folder.mkdir(parents=True, exist_ok=True)
    stem = _safe_name(filename or title) or f"document-{dt.datetime.now():%Y%m%d-%H%M%S}"
    path = None
    if destination == "word":
        try:
            path = _unique(folder / f"{stem}.docx")
            markdown_to_docx(markdown, path)
        except ImportError:
            path = None  # python-docx missing: fall back to a text file
    if path is None:
        path = _unique(folder / f"{stem}.md")
        path.write_text(markdown, encoding="utf-8")
    open_with_default_app(str(path))
    return f"Wrote '{title}' ({words} words), saved it to {path} and opened it."


def _title(markdown: str) -> str:
    m = re.search(r"^#\s+(.+)$", markdown, flags=re.M)
    return _strip_inline(m.group(1)).strip() if m else ""


def _safe_name(name: str) -> str:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", name).strip().strip(".")
    return re.sub(r"\s+", " ", name)[:80]


def _unique(path: Path) -> Path:
    n = 2
    candidate = path
    while candidate.exists():
        candidate = path.with_name(f"{path.stem} ({n}){path.suffix}")
        n += 1
    return candidate


def _strip_inline(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*|__(.+?)__", lambda m: m.group(1) or m.group(2), text)
    text = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)


def markdown_to_plain(markdown: str) -> str:
    lines = []
    for line in markdown.splitlines():
        if line.strip().startswith("```"):
            continue
        line = re.sub(r"^#{1,6}\s+", "", line)
        line = re.sub(r"^(\s*)[*+]\s+", r"\1- ", line)
        lines.append(_strip_inline(line))
    return "\n".join(lines).strip() + "\n"


def markdown_to_docx(markdown: str, path: Path) -> None:
    """Convert the simple Markdown a model writes into a tidy Word document."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    in_code = False
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            run = doc.add_paragraph().add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            continue
        if not line.strip():
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        bullet = re.match(r"^\s*[-*+]\s+(.*)", line)
        numbered = re.match(r"^\s*\d+[.)]\s+(.*)", line)
        if heading:
            level = min(len(heading.group(1)) - 1, 4)
            doc.add_heading(_strip_inline(heading.group(2)), level=level)
        elif bullet:
            _add_rich(doc.add_paragraph(style="List Bullet"), bullet.group(1))
        elif numbered:
            _add_rich(doc.add_paragraph(style="List Number"), numbered.group(1))
        else:
            _add_rich(doc.add_paragraph(), line.strip())
    doc.save(str(path))


def _add_rich(paragraph, text: str) -> None:
    """Add text to a paragraph, keeping **bold** and *italic*."""
    for token in re.split(r"(\*\*.+?\*\*|(?<!\*)\*(?!\s).+?(?<!\s)\*(?!\*))", text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            paragraph.add_run(_strip_inline(token[2:-2])).bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            paragraph.add_run(_strip_inline(token[1:-1])).italic = True
        else:
            paragraph.add_run(_strip_inline(token))

